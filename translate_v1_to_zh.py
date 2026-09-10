from pptx import Presentation

SOURCE = 'ai_talent_exchange_brainstorming_v1.pptx'
OUTPUT = 'ai_talent_exchange_brainstorming_v1_zh.pptx'
FONT_NAME = 'Microsoft YaHei'

translations = {
    'Astera Labs': 'Astera Labs',
    'AI-powered targeted intervention hub for technical solvers and small businesses.': '面向技术解决者与中小企业的 AI 精准技术干预平台。',
    'A platform that converts vague business pain into structured technical milestones, matches it to verified experts, and protects execution with escrow and AI-assisted trust.': '将模糊的业务痛点转化为结构化技术里程碑，匹配经过验证的专家，并通过托管支付与 AI 信任机制保障交付。',
    'I Need Solutions': '我需要解决方案',
    'I Provide Solutions': '我提供解决方案',
    'Why this is not a generic freelance board': '为什么这不是普通的自由职业平台',
    'The platform is positioned as a targeted technical intervention hub.': '平台定位为精准的技术干预中心。',
    'Market gap': '市场空白',
    'Small businesses need urgent, specialized fixes, not generic hiring.\nTechnical suppliers want structured, higher-trust work, not noisy gig competition.\nThe current market has too much volume and too little precision.': '中小企业需要紧急且专业的技术修复，而不是泛化招聘。\n技术供应商需要结构化、高信任度的项目，而不是嘈杂的零工竞争。\n现有市场供给很多，但精准度不足。',
    'Platform wedge': '平台切入点',
    'Convert vague pain into scoped milestones with AI translation.\nVerify supplier capability using public signals and project evidence.\nProtect execution with trust and escrow architecture.': '通过 AI 翻译将模糊痛点转化为明确的项目里程碑。\n利用公开信号和项目证据验证供应商能力。\n通过信任与托管架构保障项目执行。',
    'Strategic edge': '战略优势',
    'Specialized revenue engine, not a commodity marketplace.\nHigher-margin service layer via elastic teams and curated demand.': '专业化收入引擎，而不是同质化交易市场。\n通过弹性团队与精选需求建立更高利润的服务层。',
    'Business model and revenue engine': '商业模式与收入引擎',
    'Monetization must come from trust, matching quality, and workflow control.': '变现应来自信任、匹配质量与流程控制。',
    'Demand-side revenue': '需求侧收入',
    'Project success fee on matched interventions.\nPremium AI scoping and technical translation plans.\nEscrow facilitation and advisory uplift fees.': '按成功匹配的干预项目收取服务费。\n高级 AI 范围界定与技术翻译方案。\n托管支付服务费与顾问增值费。',
    'Supply-side revenue': '供应侧收入',
    'Proof-of-capability verification and credibility memberships.\nSubscription for visibility and inbound projects.\nRevenue share on managed squad / elastic team work.': '能力证明验证与信誉会员服务。\n提升曝光度并获得主动项目机会的订阅服务。\n托管小队与弹性团队项目的收入分成。',
    'Platform leverage': '平台杠杆',
    'AI outbound demand generation for underserved SMB niches.\nUmbrella agency model when demand exceeds supplier capacity.\nData network effects from repeated problem-to-solution loops.': '利用 AI 为服务不足的中小企业细分市场主动创造需求。\n当需求超过供应能力时，采用伞形代理与团队编排模式。\n通过反复的问题到解决方案闭环形成数据网络效应。',
    'Two-sided journey: one trust layer, two tailored experiences': '双边旅程：一套信任层，两种定制体验',
    'Demanders need clarity and privacy; suppliers need proof, signal, and qualified opportunity.': '需求方需要清晰度与隐私；供应方需要证明、信号与高质量机会。',
    'Demand / business flow': '需求方 / 企业流程',
    'Describe business pain in plain language.\nAI Technical Translator clarifies scope and anonymizes sensitive details.\nReview matched suppliers, milestones, budget, and delivery risk.\nFund smart escrow and release payment against verifiable outcomes.': '用日常语言描述业务痛点。\nAI 技术翻译器明确范围，并将敏感细节匿名化。\n查看匹配供应商、项目里程碑、预算与交付风险。\n启动智能托管，并在结果可验证后释放款项。',
    'Supply / solver flow': '供应方 / 解决者流程',
    'Connect public proof: CV, LinkedIn, GitHub, YouTube, or project pages.\nAI generates a capability profile with mapped stack and validated signals.\nReceive structured briefs and apply to high-fit interventions.\nCall out public vulnerabilities and earn funded preemptive bounties.': '连接公开证明材料：简历、LinkedIn、GitHub、YouTube 或项目页面。\nAI 生成能力档案，整理技术栈并验证相关信号。\n接收结构化需求，并申请高匹配度的技术干预项目。\n指出公开的潜在漏洞，并获得有资金支持的预防性赏金。',
    'The shared trust layer connects both journeys: privacy-safe briefs, verified capability, milestone evidence, escrow, and multidimensional reputation.': '共享信任层连接两条旅程：隐私安全的需求简报、经过验证的能力、里程碑证据、托管支付与多维信誉。',
    'What the product actually sells': '产品真正销售的是什么',
    'It is not a marketplace; it is a trusted technical execution layer.': '它不是一个市场，而是一个可信的技术执行层。',
    'For small businesses': '面向中小企业',
    'No full-time hire needed.\nFaster issue diagnosis and scoping.\nProtected project execution with escrow.\nBetter confidence in technical recommendations.': '无需招聘全职技术团队。\n更快完成问题诊断与范围界定。\n通过托管支付保护项目执行。\n更有信心地采纳技术建议。',
    'For technical solvers': '面向技术解决者',
    'Verified capability instead of generic profile noise.\nHigher-quality project matches.\nFaster access to budgeted, structured work.\nA path to larger recurring revenue via squads.': '以经过验证的能力取代泛化的个人资料噪音。\n获得更高质量的项目匹配。\n更快接触有预算且结构清晰的工作。\n通过项目小队获得更大规模的持续收入。',
    'Platform promise': '平台承诺',
    'Outcome-based technical intervention.\nTrust and AI working together.\nScalable specialized delivery.': '以结果为导向的技术干预。\n信任机制与 AI 协同工作。\n可规模化的专业交付。',
    'Logic check: where the concept is strong and where it needs tightening': '逻辑检查：概念的优势与需要加强之处',
    'The idea is compelling, but several strategic choices need sharper positioning.': '这个想法很有吸引力，但几项战略选择需要更清晰的定位。',
    'What is strong': '优势所在',
    'AI-driven scoping and translation reduce ambiguity for non-technical buyers.\nProof-of-capability improves trust and supplier quality relative to a standard gig board.\nEscrow and reputation systems create real operating confidence for both parties.\nElastic teams and preemptive pitching create a strong scale plan beyond simple matching.': 'AI 驱动的范围界定与翻译降低了非技术买家的不确定性。\n能力证明相较普通零工平台能提升信任与供应商质量。\n托管支付与信誉体系为双方建立实际的运营信心。\n弹性团队与预防性提案让平台能够超越简单匹配实现规模化。',
    'What needs improvement': '需要改进之处',
    'Define the first vertical sharply: e.g., web performance, e-commerce reliability, or security remediation.\nAvoid over-emphasizing general AI without proving real user value in the first 90 days.\nDesign the trust deposit as a quality filter, not simply a fee.\nPrevent false positives in proof-of-capability generation by mixing public-signal analysis with verified artifacts.': '清晰定义第一个垂直领域，例如网页性能、电商可靠性或安全修复。\n不要过度强调通用 AI，应先在 90 天内证明真实用户价值。\n将信任保证金设计为质量筛选机制，而不只是费用。\n通过结合公开信号分析与已验证材料，避免能力档案生成中的误判。',
    'UX foundation: role-based web platform strategy': 'UX 基础：基于角色的网页平台策略',
    'The landing page routes users into distinct demand and supplier workspaces.': '落地页将用户引导至独立的需求方与供应方工作区。',
    'Role-based entry points': '基于角色的入口',
    'Hero: clear positioning as a targeted intervention hub.\nDemand page: AI brief translation, privacy review, matching, and escrow.\nSupplier page: public-proof ingestion, capability profile, opportunities, and bounties.\nShared navigation keeps the marketplace legible without collapsing both personas.\nTrust surfaces remain visible at every conversion step.': '首屏：清晰定位为精准技术干预中心。\n需求页面：AI 简报翻译、隐私审查、匹配与托管支付。\n供应页面：公开证明材料导入、能力档案、项目机会与赏金。\n共享导航保持市场结构清晰，同时区分两类用户。\n在每个转化步骤中都展示信任信息。',
    'Web flow': '网页流程',
    'Home': '首页',
    'Demand': '需求',
    'AI scope': 'AI 范围界定',
    'Supplier': '供应方',
    'Proof': '能力证明',
    'Match + escrow': '匹配 + 托管',
    'Web platform architecture': '网页平台架构',
    'Separate role experiences sit on top of shared AI, trust, and liquidity services.': '独立的角色体验建立在共享的 AI、信任与流动性服务之上。',
    'Role experiences': '角色体验',
    'Demand workspace for problem intake and protected briefs.\nSupplier workspace for capability and opportunity discovery.\nReact + Tailwind + Framer Motion for responsive product UI.\nShared design language with role-specific information density.': '需求工作区：收集问题并创建受保护的需求简报。\n供应工作区：发现能力机会与项目机会。\n使用 React + Tailwind + Framer Motion 构建响应式产品界面。\n共享设计语言，同时根据角色调整信息密度。',
    'Shared platform services': '共享平台服务',
    'Problem and capability indexing for matching.\nVipps identity and ENK / AS entity verification status.\nMilestone escrow lifecycle management.\nEvidence, reputation, outreach review, and liquidity orchestration..': '用于匹配的问题与能力索引。\nVipps 身份验证以及 ENK / AS 企业实体验证状态。\n里程碑托管支付生命周期管理。\n证据、信誉、主动联络审查与流动性编排。',
    'AI + security layer': 'AI + 安全层',
    'Agentic scoping, anonymization, and clarification.\nPublic-signal parsing into capability profiles.\nEncrypted proprietary data and controlled disclosure.\nPredictive vulnerability detection with rate-limited safe disclosure.': '智能代理范围界定、匿名化与澄清。\n将公开信号解析为能力档案。\n加密专有数据并实施受控披露。\n预测性漏洞检测与限速安全披露。',
    'Trust and prevention: the operating model': '信任与预防：运营模型',
    'Astera reduces risk before matching, during delivery, and before a public issue becomes an avoidable incident.': 'Astera 在匹配前、交付中以及公开问题演变为可避免的事故前降低风险。',
    '1. Verify the actor': '1. 验证参与者',
    'Individual: real-name verification through Vipps.\nENK: verify proprietor and organization number.\nAS: verify authorized person and company identity.\nExpose verification status, not unnecessary identity data.': '个人：通过 Vipps 验证真实姓名。\nENK：验证经营者与组织编号。\nAS：验证授权人员与公司身份。\n展示验证状态，而不是不必要的身份数据。',
    '2. Prove the work': '2. 证明工作成果',
    'Define acceptance evidence before funds are locked.\nShow milestone state: locked, in review, or released.\nScore speed, reliability, and quality separately.\nKeep both sides informed at every handoff.': '在资金锁定前定义验收证据。\n展示里程碑状态：已锁定、审核中或已释放。\n分别评估速度、可靠性与质量。\n在每次交接时让双方保持知情。',
    '3. Prevent the incident': '3. 预防事故发生',
    'Supplier uses public signals to flag a possible client risk.\nAI drafts a bounded, respectful, non-alarmist message.\nAstera reviews and rate-limits outreach.\nAccepted signal becomes a funded diagnostic or bounty.': '供应商利用公开信号标记客户可能存在的风险。\nAI 起草范围明确、尊重客户且不过度制造恐慌的信息。\nAstera 审查并限制主动联络频率。\n被接受的风险信号转化为有资金支持的诊断或赏金项目。',
    'MVP roadmap': 'MVP 路线图',
    'Start with role-based trust workflows, then expand into predictive liquidity management.': '从基于角色的信任流程开始，再扩展到预测性流动性管理。',
    'Phase 1: demand + supplier pages': '阶段 1：需求页 + 供应页',
    'Phase 2: AI scoping + profiles': '阶段 2：AI 范围界定 + 能力档案',
    'Phase 3: escrow + reputation': '阶段 3：托管支付 + 信誉',
    'Phase 4: prevention + elastic teams': '阶段 4：预防机制 + 弹性团队',
    'The first proof point is a trusted loop: a verified person or entity submits a protected problem, receives a credible match, and releases funds against evidence. The next compounding loop lets suppliers prevent incidents before they become expensive work.': '第一个验证点是建立可信闭环：经过验证的个人或实体提交受保护的问题，获得可信匹配，并根据证据释放资金。下一层增长闭环则让供应商在事故变成高成本工作之前主动预防。',
    'Scale engine: liquidity without losing trust': '规模引擎：在不牺牲信任的前提下提升流动性',
    'The marketplace grows by actively managing imbalance between demand and supply.': '平台通过主动管理需求与供应之间的不平衡实现增长。',
    'When demand is high': '当需求旺盛时',
    'Use verified supplier profiles to assemble elastic squads.\nAdd a delivery lead and shared acceptance criteria.\nKeep escrow and reputation at the squad level.\nPrevent bottlenecks without lowering proof standards.': '利用经过验证的供应商档案组建弹性小队。\n加入交付负责人并设定共享验收标准。\n以小队为单位管理托管支付与信誉。\n在不降低证明标准的前提下避免瓶颈。',
    'When supply is high': '当供应充足时',
    'Use safe outbound campaigns to surface public risks.\nInvite clients into bounded diagnostics, not unsolicited fear.\nFund accepted prevention work through the same trust layer.\nTurn idle capability into useful, measurable intervention.': '利用安全的主动触达活动发现公开风险。\n邀请客户参与范围明确的诊断，而不是发送未经请求的恐吓信息。\n通过同一信任层为被接受的预防工作提供资金。\n将闲置能力转化为有用且可衡量的技术干预。',
    'Final strategic thesis': '最终战略论点',
    'The product becomes differentiated when AI reduces uncertainty and trust mechanisms reduce risk.': '当 AI 降低不确定性、信任机制降低风险时，产品才真正形成差异化。',
    'The platform should not be positioned as a freelance board; it should be framed as an execution infrastructure for targeted technical interventions. The most powerful combination is AI-assisted issue translation + capability verification + secure milestone-based delivery + a reputation layer that makes quality visible.': '平台不应被定位为自由职业平台，而应被定义为精准技术干预的执行基础设施。最强的组合是：AI 辅助问题翻译 + 能力验证 + 安全的里程碑交付 + 让质量可见的信誉层。',
    'If the platform can win trust early, it can scale into the operating system for SMB technical problem solving.\nTrust is not a badge added after matching. It is the system that makes prevention, matching, delivery, and scale possible.': '如果平台能够尽早赢得信任，就能发展为中小企业技术问题解决的操作系统。\n信任不是匹配之后添加的徽章，而是让预防、匹配、交付与规模化成为可能的系统。',
}


def replace_shape_text(shape):
    if not hasattr(shape, 'text_frame') or not shape.text.strip():
        return
    original = shape.text
    translated = translations.get(original)
    if translated is None:
        return

    source_lines = original.split('\n')
    translated_lines = translated.split('\n')
    paragraphs = shape.text_frame.paragraphs
    if len(paragraphs) == len(translated_lines):
        for paragraph, line in zip(paragraphs, translated_lines):
            paragraph.text = line
    else:
        shape.text_frame.text = translated

    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_NAME


presentation = Presentation(SOURCE)
for slide in presentation.slides:
    for shape in slide.shapes:
        replace_shape_text(shape)

presentation.save(OUTPUT)
print(f'Created {OUTPUT}')
